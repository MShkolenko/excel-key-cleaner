"""Generate examples/dirty-keys-demo.xlsx - a workbook to try the macro on by hand.

Every row is one kind of dirt (or one kind of cell the macro must leave alone), with the value
the macro should produce next to it. All codes are made up; they only imitate the shape of
real equipment/work keys (two letters, unit, block, tag). Needs openpyxl.

    python examples/make_demo.py
"""
import pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

OUT = pathlib.Path(__file__).with_name('dirty-keys-demo.xlsx')
NBSP, ZWSP, ZWNJ, BOM, THIN, LSEP = chr(0xA0), chr(0x200B), chr(0x200C), chr(0xFEFF), chr(0x2009), chr(0x2028)

# (what, dirty value, expected after "Выполнить" with the default boxes 1+2+4, note)
ROWS = [
    ('Кириллическая Р внутри латинского ключа', 'XX.U1.ABC.0001.10UРH.0.ZZ.YY0001-ST00', 'XX.U1.ABC.0001.10UPH.0.ZZ.YY0001-ST00', 'действие 1'),
    ('Кириллические С и Е в коротком коде', 'СЕ.U2.ABC.0002', 'CE.U2.ABC.0002', 'действие 1'),
    ('Строчная кириллическая е', 'cе.u1.abc', 'ce.u1.abc', 'действие 1: строчные а е о р с х тоже'),
    ('Русское слово - НЕ трогать', 'ВЕТЕР', 'ВЕТЕР', 'нет латиницы -> считается русским текстом'),
    ('Русская аббревиатура с цифрой - НЕ трогать', 'СМР-2', 'СМР-2', 'то же'),
    ('Кириллическая У в коде - НЕ трогать', '07UУQ', '07UУQ', 'У/у намеренно не в карте замен; попадёт в счётчик'),
    ('Пробелы по краям', '  XX.U1.ABC.0003  ', 'XX.U1.ABC.0003', 'действие 4 обрезает края'),
    ('Несколько пробелов подряд', 'XX.U1   ABC   0004', 'XX.U1 ABC 0004', 'действие 4'),
    ('Неразрывный пробел (из веба/SAP)', 'XX.U1' + NBSP + 'ABC', 'XX.U1 ABC', 'действие 4: NBSP = пробел'),
    ('Табуляция', 'XX.U1\tABC', 'XX.U1 ABC', 'действие 4'),
    ('Символы нулевой ширины', 'XX.U1' + ZWSP + '.ABC' + ZWNJ + '.0005' + BOM, 'XX.U1.ABC.0005', 'действие 4: удаляются'),
    ('Перенос строки внутри ячейки', 'XX.U1.ABC\n0006', 'XX.U1.ABC 0006', 'действие 2'),
    ('CRLF + пробелы', 'XX.U1.ABC \r\n 0007', 'XX.U1.ABC 0007', 'действия 2 и 4'),
    ('Unicode line separator', 'XX.U1.ABC' + LSEP + '0008', 'XX.U1.ABC 0008', 'действие 2'),
    ('Тонкий пробел', 'XX' + THIN + 'U1', 'XX U1', 'действие 4'),
    ('Всё сразу', '  ХХ.U1 \n ABС  ', 'XX.U1 ABC', 'действия 1, 2, 4'),
    ('Текст, похожий на число - остаётся текстом', '00123 ', '00123', 'записывается через формат "@"'),
    ('Текст, похожий на дату - остаётся текстом', '12/03 ', '12/03', 'то же'),
    ('Текст, начинающийся с "=" - НЕ станет формулой', '=1 +  1', '=1 + 1', 'то же'),
    ('Число - пропускается', 1234.5, 1234.5, 'не текст'),
    ('Формула - пропускается', '=1+1', 2, 'не константа'),
    ('#Н/Д - пропускается, не роняет макрос', '=NA()', '#Н/Д', 'ошибка в ячейке'),
    ('Пустая ячейка - пропускается', None, None, ''),
    ('Скрытая строка - пропускается', 'XX.U1.ABС.0009', 'XX.U1.ABС.0009', 'строка скрыта; попадёт в счётчик скрытых'),
    ('Разное оформление символов - пропускается', 'rich', 'rich', 'первая буква жирная; попадёт в счётчик'),
]

wb = Workbook()
ws = wb.active
ws.title = 'Демо'
ws.append(['Что проверяем', 'Значение (выдели этот столбец и запусти макрос)', 'Ожидается после 1+2+4', 'Примечание'])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='DDEBF7')
for i, (what, dirty, expected, note) in enumerate(ROWS, start=2):
    ws.cell(i, 1, what)
    cell = ws.cell(i, 2)
    exp = ws.cell(i, 3)
    ws.cell(i, 4, note)
    if what.startswith('Разное оформление'):
        cell.value = CellRichText(TextBlock(InlineFont(b=True), 'X'), 'X.U1.ABС.0011')
        exp.value = 'XX.U1.ABС.0011'
    elif isinstance(dirty, str) and dirty in ('=1+1', '=NA()'):
        cell.value = dirty                       # a real formula
        exp.value = expected
    elif dirty is None:
        exp.value = None
    elif isinstance(dirty, str):
        cell.value = dirty
        cell.data_type = 's'                     # text even when it starts with '=' - openpyxl would otherwise make it a formula
        exp.value = expected
        exp.data_type = 's'
    else:
        cell.value = dirty
        exp.value = expected
    cell.alignment = Alignment(wrap_text=False)
    if what.startswith('Скрытая'):
        ws.row_dimensions[i].hidden = True
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 44
ws.column_dimensions['C'].width = 40
ws.column_dimensions['D'].width = 48
ws.freeze_panes = 'A2'
wb.save(OUT)
print('written', OUT, len(ROWS), 'rows')
